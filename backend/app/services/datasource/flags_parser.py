"""
Flags parser — reads Datasource/Flag_Details/Server_&_Client_components_flags_details.xlsx

Two sheets:
  Sheet 1: CTCLClient(TradingStyle.txt)  → type='trading_style', source='TradingStyle.txt'
  Sheet 2: Server(CTCLManager.ini)       → type='ini_config',    source='CTCLManager.ini'

Columns used (row 1 = header, data starts row 2):
  Col B (index 1) → Trading Style Setting  → flag name
  Col C (index 2) → Usage                 → usage (Used / Not in use)
  Col D (index 3) → Possible values       → possible values
  Col E (index 4) → Description           → description
  All other columns ignored.
"""

import logging
from pathlib import Path
from typing import Any, Dict, List

from app.services.datasource.base_parser import BaseParser

logger = logging.getLogger(__name__)

# Map sheet name → type and source labels
SHEET_META = {
    "CTCLClient(TradingStyle.txt)": {
        "type": "trading_style",
        "source": "TradingStyle.txt",
        "module": "CLIENT",
        "section": "CTCL_CLIENT",
    },
    "Server(CTCLManager.ini)": {
        "type": "ini_config",
        "source": "CTCLManager.ini",
        "module": "SERVER",
        "section": "CTCL_SERVER",
    },
}


class FlagsParser(BaseParser):
    """Parses trading style and INI config flags from the xlsx flag details file."""

    def get_source_name(self) -> str:
        return "Flags"

    def parse(self) -> List[Dict[str, Any]]:
        flags: List[Dict[str, Any]] = []
        folder = self.root / "Flag_Details"
        if not folder.exists():
            logger.warning("FlagsParser: Flag_Details/ folder not found")
            return []

        for xlsx_file in folder.glob("*.xlsx"):
            flags.extend(self._parse_xlsx(xlsx_file))

        logger.info(
            f"FlagsParser: parsed {len(flags)} flags from Flag_Details/ "
            f"(trading_style={sum(1 for f in flags if f['type']=='trading_style')}, "
            f"ini_config={sum(1 for f in flags if f['type']=='ini_config')})"
        )
        return flags

    def _parse_xlsx(self, path: Path) -> List[Dict[str, Any]]:
        flags: List[Dict[str, Any]] = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

            for sheet_name in wb.sheetnames:
                meta = SHEET_META.get(sheet_name)
                if meta is None:
                    # Try partial match
                    for key, val in SHEET_META.items():
                        if key.lower() in sheet_name.lower() or sheet_name.lower() in key.lower():
                            meta = val
                            break
                if meta is None:
                    logger.debug(f"FlagsParser: unknown sheet '{sheet_name}', skipping")
                    continue

                ws = wb[sheet_name]
                rows = list(ws.iter_rows(min_row=2, values_only=True))  # skip title row 0

                for idx, row in enumerate(rows):
                    # Skip completely empty rows
                    if not row or not any(row):
                        continue
                    # Col A = Sr No (index 0), Col B = name (index 1)
                    name_raw = row[1] if len(row) > 1 else None
                    if name_raw is None:
                        continue
                    name = str(name_raw).strip()
                    if not name or name.lower() in ("trading style setting", "sr no.", ""):
                        continue  # Skip header-like rows

                    usage = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
                    possible_values = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
                    description = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""

                    flags.append({
                        "id": f"{meta['section']}_{idx}",
                        "name": name,
                        "value": possible_values or "-",
                        "type": meta["type"],
                        "section": meta["section"],
                        "module": meta["module"],
                        "source_file": meta["source"],
                        "usage": usage,
                        "possible_values": possible_values,
                        "description": description[:500] if description else f"{name} flag",
                        "is_active": usage.lower() == "used",
                    })

            wb.close()
        except Exception as e:
            logger.error(f"FlagsParser xlsx error for {path.name}: {e}", exc_info=True)

        return flags
