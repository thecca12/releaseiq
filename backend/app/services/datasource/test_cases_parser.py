"""
Test Cases parser — reads Datasource/Test_Cases/Ajay_Team_Test_Cases_.xlsx

Structure:
  Sheet 1 "All Modules" — index: Sr no | Modules | Tester Name
  Sheets 2-84 — each is one module with format:
    Row 0:   "All Module" / "All Modules"
    Row 1-2: "Module Name : <name>"
    Row 2-3: "Tester Name : <tester>"
    Row 3-4: Headers → Test case ID | Test case Description | Created By | Reviewed By | module name | Test Steps
    Row 4+:  Actual test case data

Columns extracted:
    0 (A) → test_id
    1 (B) → description
    2 (C) → created_by
    3 (D) → reviewed_by
    4 (E) → sub_module
    5 (F) → steps
"""
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from app.services.datasource.base_parser import BaseParser

logger = logging.getLogger(__name__)


class TestCasesParser(BaseParser):

    def get_source_name(self) -> str:
        return "Test Cases"

    def parse(self) -> List[Dict[str, Any]]:
        folder = self.root / "Test_Cases"
        if not folder.exists():
            logger.warning("TestCasesParser: Test_Cases/ folder not found")
            return []

        for xlsx_file in folder.glob("*.xlsx"):
            result = self._parse_xlsx(xlsx_file)
            if result:
                logger.info(f"TestCasesParser: parsed {len(result)} test cases from {xlsx_file.name} ({len(set(r['module'] for r in result))} modules)")
                return result

        return []

    # ── Core xlsx parser ──────────────────────────────────────────────────────

    def _parse_xlsx(self, path: Path) -> List[Dict[str, Any]]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            logger.error(f"TestCasesParser: cannot open {path.name}: {e}")
            return []

        all_cases: List[Dict[str, Any]] = []

        # Skip the first sheet ("All Modules") — it's just the index
        module_sheets = [s for s in wb.sheetnames if s.strip().lower() not in ("all modules", "all module")]

        for sheet_name in module_sheets:
            try:
                ws = wb[sheet_name]
                cases = self._parse_sheet(ws, sheet_name)
                all_cases.extend(cases)
            except Exception as e:
                logger.debug(f"TestCasesParser: skipping sheet {sheet_name!r}: {e}")

        wb.close()
        return all_cases

    def _parse_sheet(self, ws, sheet_name: str) -> List[Dict[str, Any]]:
        rows = list(ws.iter_rows(max_row=500, values_only=True))
        if not rows:
            return []

        # Extract module name and tester from header rows
        module_name = sheet_name.strip()
        tester_name = ""

        for row in rows[:6]:
            if not row:
                continue
            cell0 = str(row[0] or "").strip()
            if cell0.lower().startswith("module name"):
                extracted = cell0.split(":", 1)[-1].strip() if ":" in cell0 else ""
                if extracted:
                    module_name = extracted
            elif cell0.lower().startswith("tester name"):
                tester_name = cell0.split(":", 1)[-1].strip() if ":" in cell0 else ""

        # Find the column header row (row containing "Test case ID")
        header_row_idx = -1
        for i, row in enumerate(rows[:8]):
            if not row:
                continue
            row_text = " ".join(str(v or "").lower() for v in row)
            if "test case id" in row_text or "test case description" in row_text:
                header_row_idx = i
                break

        if header_row_idx == -1:
            return []

        # Parse test cases from rows after the header
        cases: List[Dict[str, Any]] = []
        for row in rows[header_row_idx + 1:]:
            if not row or not row[0]:
                continue

            # Skip sub-header rows that repeat the module name
            cell0 = str(row[0] or "").strip()
            if not cell0 or cell0.lower() in ("all modules", "all module", "module name", "test case id"):
                continue
            # Must be a numeric ID
            try:
                int(str(cell0).split(".")[0])
            except ValueError:
                continue

            description = str(row[1] or "").strip() if len(row) > 1 and row[1] else ""
            created_by  = str(row[2] or "").strip() if len(row) > 2 and row[2] else ""
            reviewed_by = str(row[3] or "").strip() if len(row) > 3 and row[3] else ""
            sub_module  = str(row[4] or "").strip() if len(row) > 4 and row[4] else ""
            steps       = str(row[5] or "").strip() if len(row) > 5 and row[5] else ""

            if not description:
                continue

            cases.append({
                "test_id":     f"{sheet_name[:20]}-{cell0}",
                "test_name":   description[:200],
                "module":      module_name,
                "sub_module":  sub_module or module_name,
                "sheet":       sheet_name,
                "tester":      tester_name,
                "type":        "Functional",
                "steps":       steps[:500],
                "created_by":  created_by,
                "reviewed_by": reviewed_by,
                "expected_result": "Pass",
                "status":      "Active",
                "priority":    "Medium",
                "last_run":    "",
                "automation":  "Manual",
            })

        return cases
